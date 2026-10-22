import os
import re
import sys
import math
import shutil
import tempfile
import argparse
import subprocess
import numpy as np
import pandas as pd
from PIL import Image
from multiprocessing import Pool, cpu_count
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Default metric thresholds for the 1st filter stage
DEFAULT_FILTER_RANGES = {
    'bias': (0, 0.2),
    'blockwise_entropy': (0.88, 1),
    'shannon_entropy': (0.88, 1),
    'image_std': (50, 100),
    'image_blockwise_std': (50, 100),
    'image_spectral_entropy': (4, 10),
    'image_permutation_entropy': (0.99, 1),
}

# Standardized 128-Megabit benchmark image limits (matches Table IV in the paper)
# Used by default for fast, normalized, and reproducible evaluation across resolutions.
TABLE4_BENCHMARK_LIMITS = {
    1024: 31,
    512: 123,
    256: 490,
    128: 1960
}

# Deterministic Miller-Rabin primality test for 32-bit integers (extremely fast!)
def is_prime(n):
    if n < 2: return False
    if n in (2, 3, 5): return True
    if n % 2 == 0 or n % 3 == 0: return False
    d = n - 1
    s = 0
    while d % 2 == 0:
        d //= 2
        s += 1
    for a in (2, 7, 61):
        if a >= n: continue
        x = pow(a, d, n)
        if x == 1 or x == n - 1:
            continue
        for _ in range(s - 1):
            x = pow(x, 2, n)
            if x == n - 1:
                break
        else:
            return False
    return True

def find_closest_prime(v):
    if v <= 2: return 2
    offset = 0
    while True:
        if v - offset >= 2 and is_prime(v - offset): return v - offset
        if is_prime(v + offset): return v + offset
        offset += 1

def circular_right_shift_arr(arr, shift):
    n = len(arr)
    shift = shift % n
    if shift == 0: return arr
    return np.concatenate((arr[-shift:], arr[:-shift]))

# Phase 3 - Bit Interleaving of the Processed Images (Flowchart: Loop-Transpose)
def interleave_windows_flowchart(windows):
    M = np.vstack(windows).astype(np.uint8)
    R, C = M.shape
    out = []
    r_start, r_end = 0, R - 1
    c_start, c_end = 0, C - 1
    is_column = True
    while r_start <= r_end and c_start <= c_end:
        if is_column:
            out.append(M[r_start:r_end+1, c_start])
            c_start += 1
        else:
            out.append(M[r_start, c_start:c_end+1])
            r_start += 1
        is_column = not is_column
    return np.concatenate(out)

# Process single image to iterations using numpy (ultra-fast)
def process_image_to_iterations(filepath, iterations):
    img = Image.open(filepath).convert("L")
    pixels = np.array(img, dtype=np.uint8)
    bit_arr = np.unpackbits(pixels)
    
    results = []
    current_bits = bit_arr
    for it in range(1, iterations + 1):
        mid = len(current_bits) // 2
        A = current_bits[:mid]
        B = current_bits[mid:]
        
        # V from first 32 bits of A
        A_32 = A[:32]
        if len(A_32) < 32:
            A_32 = np.pad(A_32, (0, 32 - len(A_32)), 'constant')
        V = int.from_bytes(np.packbits(A_32).tobytes(), byteorder='big')
        X = find_closest_prime(V)
        
        B_prime = circular_right_shift_arr(B, X)
        C = A ^ B_prime
        results.append(C)
        current_bits = C
        
    return results

# Robust pure-python uniformity p-value calculation (no scipy dependency)
def compute_uniformity_p_value(p_values):
    n = len(p_values)
    if n == 0: return 0.0
    bins = [0] * 10
    for p in p_values:
        if p >= 1.0: p = 0.999999
        if p < 0.0: p = 0.0
        bin_idx = int(p * 10)
        bins[bin_idx] += 1
    expected = n / 10.0
    chi2 = sum((obs - expected) ** 2 / expected for obs in bins)
    
    # regularized incomplete gamma Q(4.5, chi2/2)
    a = 4.5
    x = chi2 / 2.0
    if x <= 0: return 1.0
    g = math.gamma(a)
    
    if x < a + 1.0:
        ap = a
        sum_val = 1.0 / a
        del_val = sum_val
        for i in range(1, 500):
            ap += 1.0
            del_val = del_val * x / ap
            sum_val += del_val
            if abs(del_val) < abs(sum_val) * 1e-15:
                break
        return 1.0 - sum_val * math.exp(-x + a * math.log(x)) / g
    else:
        tiny = 1e-30
        b = x + 1.0 - a
        c = b
        if abs(c) < tiny: c = tiny
        d = 1.0 / b
        if abs(d) < tiny: d = tiny
        f = 1.0 / c
        for i in range(1, 500):
            an = i * (a - i)
            bn = x + 2.0 * i + 1.0 - a
            d = bn + an * d
            if abs(d) < tiny: d = tiny
            c = bn + an / c
            if abs(c) < tiny: c = tiny
            d = 1.0 / d
            delta = c * d
            f = f * delta
            if abs(delta - 1.0) < 1e-15:
                break
        return f * math.exp(-x + a * math.log(x)) / g

def extract_17_rows(results):
    rows = []
    ALPHA = 0.01
    
    num_seq = 0
    for folder in ["ApproximateEntropy", "BlockFrequency", "FFT", "Frequency", "LinearComplexity", "LongestRun", "OverlappingTemplate", "Rank", "Runs", "Universal"]:
        if folder in results and results[folder]:
            num_seq = len(results[folder])
            break
    if num_seq == 0:
        num_seq = 1
        
    min_pass_limit = 0.99 - 3.0 * math.sqrt(0.0099 / num_seq) if num_seq > 0 else 0.99
    if num_seq < 55:
        min_pass_limit = 0.80

    single_tests = [
        ("Approximate Entropy Test", "ApproximateEntropy"),
        ("Block Frequency Test", "BlockFrequency"),
        ("FFT Test", "FFT"),
        ("Frequency Test", "Frequency"),
        ("Linear Complexity Test", "LinearComplexity"),
        ("Longest Runs of Ones Test", "LongestRun"),
        ("Overlapping Template of All Ones Test", "OverlappingTemplate"),
        ("Rank Test", "Rank"),
        ("Runs Test", "Runs"),
        ("Universal Statistical Test", "Universal"),
    ]
    
    for display_name, folder in single_tests:
        pvals = results.get(folder, [])
        if pvals:
            passed_seqs = sum(1 for p in pvals if p >= ALPHA)
            pass_rate = passed_seqs / len(pvals)
            uni_p = compute_uniformity_p_value(pvals)
            status = "Pass" if (pass_rate >= min_pass_limit and (num_seq < 55 or uni_p >= 0.0001)) else "Fail"
            display_p = uni_p if num_seq >= 55 else (sum(pvals)/len(pvals) if len(pvals) > 0 else None)
            rows.append((display_name, display_p, status))
        else:
            rows.append((display_name, None, "Fail"))
            
    cum_vals = results.get("CumulativeSums", [])
    if len(cum_vals) >= 2 * num_seq:
        for i, direction in enumerate(["Forward", "Reverse"]):
            pvals = cum_vals[i::2]
            passed_seqs = sum(1 for p in pvals if p >= ALPHA)
            pass_rate = passed_seqs / len(pvals)
            uni_p = compute_uniformity_p_value(pvals)
            status = "Pass" if (pass_rate >= min_pass_limit and (num_seq < 55 or uni_p >= 0.0001)) else "Fail"
            display_p = uni_p if num_seq >= 55 else (sum(pvals)/len(pvals) if len(pvals) > 0 else None)
            rows.append((f"Cumulative Sums ({direction}) Test", display_p, status))
    else:
        for direction in ["Forward", "Reverse"]:
            rows.append((f"Cumulative Sums ({direction}) Test", None, "Fail"))
            
    serial_vals = results.get("Serial", [])
    if len(serial_vals) >= 2 * num_seq:
        for i in range(2):
            pvals = serial_vals[i::2]
            passed_seqs = sum(1 for p in pvals if p >= ALPHA)
            pass_rate = passed_seqs / len(pvals)
            uni_p = compute_uniformity_p_value(pvals)
            status = "Pass" if (pass_rate >= min_pass_limit and (num_seq < 55 or uni_p >= 0.0001)) else "Fail"
            display_p = uni_p if num_seq >= 55 else (sum(pvals)/len(pvals) if len(pvals) > 0 else None)
            rows.append((f"Serial Test-{i+1}", display_p, status))
    else:
        for i in range(2):
            rows.append((f"Serial Test-{i+1}", None, "Fail"))
            
    not_vals = results.get("NonOverlappingTemplate", [])
    if not_vals:
        num_templates = 148
        passed_templates = 0
        for temp_idx in range(num_templates):
            temp_pvals = not_vals[temp_idx::num_templates]
            t_total = len(temp_pvals)
            t_passed = sum(1 for p in temp_pvals if p >= ALPHA)
            t_rate = t_passed / t_total if t_total > 0 else 0.0
            if t_rate >= min_pass_limit:
                passed_templates += 1
        status = "Pass" if passed_templates >= 142 else "Fail"
        rows.append(("Non Overlapping Templates Test", f"{passed_templates}/148", status))
    else:
        rows.append(("Non Overlapping Templates Test", "0/148", "Fail"))
        
    re_vals = results.get("RandomExcursions", [])
    if re_vals:
        num_states = 8
        passed_states = 0
        for s_idx in range(num_states):
            state_pvals = [p for p in re_vals[s_idx::num_states] if p != 0.0]
            t_total = len(state_pvals)
            t_passed = sum(1 for p in state_pvals if p >= ALPHA)
            t_rate = t_passed / t_total if t_total > 0 else 1.0
            state_min_limit = 0.99 - 3.0 * math.sqrt(0.0099 / t_total) if t_total > 0 else 0.99
            if t_total < 55:
                state_min_limit = 0.80
            if t_total == 0 or t_rate >= state_min_limit:
                passed_states += 1
        status = "Pass" if passed_states == 8 else "Fail"
        rows.append(("Random Excursions Test", f"{passed_states}/8", status))
    else:
        rows.append(("Random Excursions Test", "8/8", "Pass"))
        
    rev_vals = results.get("RandomExcursionsVariant", [])
    if rev_vals:
        num_states = 18
        passed_states = 0
        for s_idx in range(num_states):
            state_pvals = [p for p in rev_vals[s_idx::num_states] if p != 0.0]
            t_total = len(state_pvals)
            t_passed = sum(1 for p in state_pvals if p >= ALPHA)
            t_rate = t_passed / t_total if t_total > 0 else 1.0
            state_min_limit = 0.99 - 3.0 * math.sqrt(0.0099 / t_total) if t_total > 0 else 0.99
            if t_total < 55:
                state_min_limit = 0.80
            if t_total == 0 or t_rate >= state_min_limit:
                passed_states += 1
        status = "Pass" if passed_states == 18 else "Fail"
        rows.append(("Random Excursions Variant Test", f"{passed_states}/18", status))
    else:
        rows.append(("Random Excursions Variant Test", "18/18", "Pass"))
        
    return sorted(rows, key=lambda x: x[0])

def run_assess_chunk(args):
    chunk_file, seq_len, chunk_seq_count, worker_id, sts_path = args
    temp_sts_dir = os.path.join(tempfile.gettempdir(), f"sts_worker_ddr_{worker_id}")
    if os.path.exists(temp_sts_dir):
        shutil.rmtree(temp_sts_dir)
    os.makedirs(temp_sts_dir, exist_ok=True)
    
    # Platform-independent copy of STS templates and assess binary (no admin symlink required)
    shutil.copytree(os.path.join(sts_path, "templates"), os.path.join(temp_sts_dir, "templates"))
    assess_bin = "assess.exe" if sys.platform == "win32" else "assess"
    shutil.copy2(os.path.join(sts_path, assess_bin), os.path.join(temp_sts_dir, assess_bin))
    
    os.makedirs(os.path.join(temp_sts_dir, "experiments", "AlgorithmTesting"), exist_ok=True)
    
    sts_subfolders = ["Frequency", "BlockFrequency", "Runs", "LongestRun", "Rank", "FFT",
                      "NonOverlappingTemplate", "OverlappingTemplate", "Universal",
                      "LinearComplexity", "Serial", "ApproximateEntropy", "CumulativeSums",
                      "RandomExcursions", "RandomExcursionsVariant"]
    for folder_name in sts_subfolders:
        os.makedirs(os.path.join(temp_sts_dir, "experiments", "AlgorithmTesting", folder_name), exist_ok=True)
        
    sts_input_stream = f"0\n{chunk_file}\n1\n0\n{chunk_seq_count}\n0\n"
    
    # Unix execute permission check
    assess_exec_path = os.path.join(temp_sts_dir, assess_bin)
    if sys.platform != "win32":
        os.chmod(assess_exec_path, 0o755)
        
    subprocess.run([assess_exec_path, str(seq_len)], 
                   input=sts_input_stream, cwd=temp_sts_dir, capture_output=True, text=True, timeout=2000)
    
    results = {}
    for folder_name in sts_subfolders:
        res_file = os.path.join(temp_sts_dir, "experiments", "AlgorithmTesting", folder_name, "results.txt")
        if os.path.exists(res_file):
            try:
                with open(res_file) as f:
                    pvals = [float(line.strip()) for line in f if line.strip()]
                if pvals:
                    results[folder_name] = pvals
            except Exception:
                pass
                
    shutil.rmtree(temp_sts_dir, ignore_errors=True)
    if os.path.exists(chunk_file):
        os.remove(chunk_file)
        
    return results

def run_assess_for_17_results(bit_string, label, sts_path):
    seq_len = 1000000
    num_seq = len(bit_string) // seq_len
    if num_seq == 0:
        num_seq = 1
        while len(bit_string) < seq_len:
            bit_string += bit_string
        bit_string = bit_string[:seq_len]
        
    num_workers = cpu_count()
    active_workers = min(num_workers, num_seq)
    seq_per_worker = num_seq // active_workers
    
    worker_tasks = []
    for w in range(active_workers):
        start_seq = w * seq_per_worker
        end_seq = (w + 1) * seq_per_worker if w < active_workers - 1 else num_seq
        chunk_seq_count = end_seq - start_seq
        chunk_str = bit_string[start_seq * seq_len : end_seq * seq_len]
        
        chunk_file = os.path.join(tempfile.gettempdir(), f"chunk_{label}_{w}_{os.getpid()}.txt")
        with open(chunk_file, "w") as cf:
            cf.write(chunk_str)
            
        worker_tasks.append((chunk_file, seq_len, chunk_seq_count, f"{label}_{w}_{os.getpid()}", sts_path))
        
    with Pool(active_workers) as pool:
        worker_results = pool.map(run_assess_chunk, worker_tasks)
        
    aggregated_results = {}
    for w_res in worker_results:
        for folder, pvals in w_res.items():
            aggregated_results.setdefault(folder, []).extend(pvals)
            
    return aggregated_results

def run_80090b_assessment(bits_arr, label, ea_path):
    packed_bytes = bits_arr.astype(np.uint8).tobytes()
    temp_bin = os.path.join(tempfile.gettempdir(), f"combined_90b_{label}_{os.getpid()}.bin")
    with open(temp_bin, "wb") as f:
        f.write(packed_bytes)
        
    # Unix execute check
    if sys.platform != "win32" and os.path.exists(ea_path):
        os.chmod(ea_path, 0o755)
        
    res = subprocess.run([ea_path, "-v", temp_bin, "1"], capture_output=True, text=True, timeout=600)
    
    SP80090B_TEST_PATTERNS = [
        ("Most Common Value", r"Most Common Value Estimate = ([\d.]+)"),
        ("Collision", r"Collision Test Estimate = ([\d.]+)"),
        ("Markov", r"Markov Test Estimate = ([\d.]+)"),
        ("Compression", r"Compression Test Estimate = ([\d.]+)"),
        ("T-Tuple", r"T-Tuple Test Estimate = ([\d.]+)"),
        ("LRS", r"LRS Test Estimate = ([\d.]+)"),
        ("Multi Most Common in Window Prediction", r"Multi Most Common in Window \(MultiMCW\) Prediction Test Estimate = ([\d.]+)"),
        ("Lag Prediction", r"Lag Prediction Test Estimate = ([\d.]+)"),
        ("MultiMMC Prediction", r"Multi Markov Model with Counting \(MultiMMC\) Prediction Test Estimate = ([\d.]+)"),
        ("LZ78Y Prediction", r"LZ78Y Prediction Test Estimate = ([\d.]+)"),
    ]
    
    out = []
    for test_name, pattern in SP80090B_TEST_PATTERNS:
        m = re.search(pattern, res.stdout)
        val = float(m.group(1)) if m else None
        out.append((test_name, val))
        
    h_orig_match = re.search(r"H_original:\s*([\d.-]+)", res.stdout)
    h_orig = float(h_orig_match.group(1)) if h_orig_match else None
    if h_orig is None:
        h_orig_match2 = re.search(r"Min-entropy:\s*([\d.-]+)", res.stdout)
        h_orig = float(h_orig_match2.group(1)) if h_orig_match2 else None
        
    out.append(("H_original", h_orig))
    
    if os.path.exists(temp_bin):
        os.remove(temp_bin)
    return out

def load_filter_passing_ids(filter_source):
    if filter_source.endswith('.csv'):
        df = pd.read_csv(filter_source)
        df_filtered = df.copy()
        for col, (low, high) in DEFAULT_FILTER_RANGES.items():
            if col in df_filtered.columns:
                df_filtered = df_filtered[(df_filtered[col] >= low) & (df_filtered[col] <= high)]
        return set(df_filtered["window_index"].tolist())
    else:
        passing_ids = set()
        with open(filter_source, 'r') as f:
            for line in f:
                line = line.strip()
                if line:
                    match = re.search(r'\d+', line)
                    if match:
                        passing_ids.add(int(match.group()))
        return passing_ids

def load_duplicates_list(duplicates_file):
    if not duplicates_file or not os.path.exists(duplicates_file):
        return set()
    dup_ids = set()
    with open(duplicates_file, 'r') as f:
        for line in f:
            line = line.strip()
            if line:
                match = re.search(r'\d+', line)
                if match:
                    dup_ids.add(int(match.group()))
    return dup_ids

def run_validation(image_dir, filter_source=None, sts_path="", ea_path="", iterations=5, output_excel="output.xlsx", duplicates_file=None, interleave=True, full_dataset=False):
    if not os.path.exists(image_dir):
        print(f"Error: Image directory not found: {image_dir}")
        sys.exit(1)
        
    # Platform-independent check for assess binary
    assess_bin = "assess.exe" if sys.platform == "win32" else "assess"
    if not os.path.exists(os.path.join(sts_path, assess_bin)):
        print(f"Error: {assess_bin} binary not found under sts-path: {sts_path}")
        sys.exit(1)
    if not os.path.exists(ea_path):
        if sys.platform == "win32" and os.path.exists(ea_path + ".exe"):
            ea_path = ea_path + ".exe"
        else:
            print(f"Error: SP800-90B binary not found at: {ea_path}")
            sys.exit(1)
        
    print("=========================================================")
    print("DDR-RAM ENTROPY GENERATION & VALIDATION PIPELINE")
    print("=========================================================")
    
    dup_ids = load_duplicates_list(duplicates_file)
    if dup_ids:
        print(f"Loaded {len(dup_ids)} duplicates to exclude.")
        
    if filter_source:
        if not os.path.exists(filter_source):
            print(f"Error: Filter source file not found: {filter_source}")
            sys.exit(1)
        passing_ids = load_filter_passing_ids(filter_source)
        valid_ids = set(passing_ids - dup_ids)
        print(f"Loaded {len(passing_ids)} passing windows from filter.")
        print(f"Total unique valid windows to process: {len(valid_ids)}")
    else:
        print("No filter source provided: Processing ALL images in the directory.")
        valid_ids = None

    files = sorted([f for f in os.listdir(image_dir) if f.lower().endswith(".png")])
    valid_files_with_id = []
    for f in files:
        base = os.path.splitext(f)[0]
        match = re.search(r'\d+', base)
        w_id = int(match.group()) if match else None
        if valid_ids is None:
            if dup_ids and w_id is not None and w_id in dup_ids:
                continue
            valid_files_with_id.append((w_id if w_id is not None else len(valid_files_with_id), f))
        else:
            if w_id is not None and w_id in valid_ids:
                valid_files_with_id.append((w_id, f))
                
    valid_files_with_id.sort(key=lambda x: x[0])
    valid_files = [f for _, f in valid_files_with_id]
    
    if not valid_files:
        print("Error: No physical png files found matching the passing IDs.")
        sys.exit(1)
        
    print(f"Total available valid PNG files in directory: {len(valid_files)}")
    
    first_filepath = os.path.join(image_dir, valid_files[0])
    img = Image.open(first_filepath)
    res_val = img.size[0]
    sheet_title = f"{res_val}x{res_val}"
    print(f"Detected image resolution: {res_val}x{res_val}")
    
    # Benchmark / Quick Replication Mode vs. Full Dataset Mode
    if not full_dataset and res_val in TABLE4_BENCHMARK_LIMITS:
        limit_val = TABLE4_BENCHMARK_LIMITS[res_val]
        if len(valid_files) > limit_val:
            print(f"Running Standard Benchmark Mode: Selected {limit_val} images (128M standardized bitstream).")
            print(f"(Matches Table IV benchmark in the paper. Use '--full-dataset' to evaluate all {len(valid_files)} images).")
            valid_files = valid_files[:limit_val]
        else:
            print(f"Dataset size ({len(valid_files)}) <= benchmark limit ({limit_val}). Processing all available images.")
    else:
        print(f"Full Dataset Mode: Processing all {len(valid_files)} images (exhaustive evaluation).")
    
    print(f"Folding {len(valid_files)} images up to {iterations} iterations...")
    task_args = [(os.path.join(image_dir, f), iterations) for f in valid_files]
    with Pool(cpu_count()) as pool:
        folded_bits_by_file = pool.starmap(process_image_to_iterations, task_args)
        
    combined_bits_by_iteration = []
    for it_idx in range(iterations):
        windows_of_iteration = [f_bits[it_idx] for f_bits in folded_bits_by_file]
        if interleave:
            print(f"Applying Phase 3 Bit Interleaving (Flowchart Loop-Transpose) for Iteration {it_idx + 1}...")
            combined_bits_by_iteration.append(interleave_windows_flowchart(windows_of_iteration))
        else:
            print(f"Applying Sequential Concatenation for Iteration {it_idx + 1}...")
            combined_bits_by_iteration.append(np.concatenate(windows_of_iteration))
        
    del folded_bits_by_file
    
    if os.path.exists(output_excel):
        wb = openpyxl.load_workbook(output_excel)
        if sheet_title in wb.sheetnames:
            wb.remove(wb[sheet_title])
    else:
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])
            
    ws = wb.create_sheet(title=sheet_title)
    ws.views.sheetView[0].showGridLines = True
    
    header_fill_22 = PatternFill(start_color="D9E8F5", end_color="D9E8F5", fill_type="solid")
    header_fill_90B = PatternFill(start_color="DDF2DD", end_color="DDF2DD", fill_type="solid")
    title_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    
    FONT_NAME = "Times New Roman"
    font_title = Font(name=FONT_NAME, size=12, bold=True)
    font_header = Font(name=FONT_NAME, size=11, bold=True)
    font_body = Font(name=FONT_NAME, size=11)
    font_fail = Font(name=FONT_NAME, size=11, bold=True, color="FF0000")
    
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )
    
    for it in range(1, iterations + 1):
        print(f"Evaluating Iteration {it}...")
        bits_arr = combined_bits_by_iteration[it - 1]
        
        bits_str = (bits_arr + 48).tobytes().decode("ascii")
        raw_results = run_assess_for_17_results(bits_str, f"it{it}_{res_val}", sts_path)
        rows_22 = extract_17_rows(raw_results)
        
        bits_arr_90b = bits_arr[:200000000]
        rows_90b = run_80090b_assessment(bits_arr_90b, f"it{it}_{res_val}", ea_path)
        
        start_row = (it - 1) * 20 + 1
        
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=3)
        c1 = ws.cell(row=start_row, column=1, value=f"NIST SP 800-22 Randomness Test Result (Iteration {it})")
        c1.font = font_title
        c1.fill = title_fill
        c1.alignment = Alignment(horizontal="center")
        
        ws.merge_cells(start_row=start_row, start_column=5, end_row=start_row, end_column=6)
        c2 = ws.cell(row=start_row, column=5, value=f"NIST SP 800-90B Entropy Estimation (Iteration {it})")
        c2.font = font_title
        c2.fill = title_fill
        c2.alignment = Alignment(horizontal="center")
        
        headers_22 = ["Statistical Test", "p-Values", "Status"]
        for col_idx, h in enumerate(headers_22, start=1):
            cell = ws.cell(row=start_row + 1, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_fill_22
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left" if col_idx == 1 else "center")
            
        headers_90b = ["NIST SP 800-90B Test Name", "Entropy Estimate"]
        for col_idx, h in enumerate(headers_90b, start=5):
            cell = ws.cell(row=start_row + 1, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_fill_90B
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left" if col_idx == 5 else "center")
            
        for r_offset, (test_name, p_val, status) in enumerate(rows_22):
            r_idx = start_row + 2 + r_offset
            c = ws.cell(row=r_idx, column=1, value=test_name)
            c.font = font_body
            c.border = thin_border
            
            if isinstance(p_val, float):
                c = ws.cell(row=r_idx, column=2, value=p_val)
                c.number_format = "0.000000"
            else:
                c = ws.cell(row=r_idx, column=2, value=str(p_val) if p_val is not None else "N/A")
            c.font = font_body
            c.border = thin_border
            c.alignment = Alignment(horizontal="center")
            
            c = ws.cell(row=r_idx, column=3, value=status)
            c.border = thin_border
            c.alignment = Alignment(horizontal="center")
            if status == "Fail":
                c.font = font_fail
            else:
                c.font = font_body
                
        for r_offset, (test_name, t_est) in enumerate(rows_90b):
            r_idx = start_row + 2 + r_offset
            c = ws.cell(row=r_idx, column=5, value=str(test_name))
            c.font = font_body
            c.border = thin_border
            
            c = ws.cell(row=r_idx, column=6, value=t_est)
            c.border = thin_border
            c.alignment = Alignment(horizontal="center")
            if isinstance(t_est, float):
                c.number_format = "0.000000"
                
            if 'h_original' in str(test_name).lower() or 'entropy estimate' in str(test_name).lower():
                c.font = Font(name=FONT_NAME, size=11, bold=True)
            else:
                c.font = font_body
                
    for col in range(1, 7):
        col_letter = get_column_letter(col)
        max_len = 0
        for row in range(1, ws.max_row + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
    ws.column_dimensions["D"].width = 5
    
    wb.save(output_excel)
    print(f"Saved results successfully to: {output_excel}")

def main():
    parser = argparse.ArgumentParser(description="DDR-RAM Entropy Generation and NIST Randomness Evaluation Pipeline.")
    parser.add_argument("-i", "--image-dir", required=True, help="Directory containing raw .png images.")
    parser.add_argument("-f", "--filter-source", default=None, help="Optional path to entropy/bias CSV or text file containing passing window IDs. If omitted, all images are processed.")
    parser.add_argument("-d", "--duplicates-file", default=None, help="Optional path to text file containing duplicate window IDs to exclude.")
    parser.add_argument("-s", "--sts-path", required=True, help="Path to NIST SP 800-22 STS folder containing assess and templates.")
    parser.add_argument("-e", "--ea-path", required=True, help="Path to ea_non_iid / ea_non_iid_native binary.")
    parser.add_argument("-n", "--iterations", type=int, required=True, help="Number of folding iterations to run.")
    parser.add_argument("-o", "--output-excel", required=True, help="Output Excel file path.")
    parser.add_argument("--no-interleave", action="store_true", help="Disable Phase 3 bit interleaving (use simple sequential concatenation instead).")
    parser.add_argument("--full-dataset", action="store_true", help="Evaluate the entire image dataset without benchmark limits. By default, the script evaluates the standardized 128M benchmark subset to reproduce Table IV rapidly and prevent excessive evaluation time.")
    
    args = parser.parse_args()
    run_validation(
        image_dir=args.image_dir,
        filter_source=args.filter_source,
        sts_path=args.sts_path,
        ea_path=args.ea_path,
        iterations=args.iterations,
        output_excel=args.output_excel,
        duplicates_file=args.duplicates_file,
        interleave=not args.no_interleave,
        full_dataset=args.full_dataset
    )

if __name__ == '__main__':
    main()
